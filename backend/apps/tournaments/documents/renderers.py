"""Renderizadores PDF (ReportLab) y DOCX (python-docx) de bloques de documento."""
import io
from html import escape

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import PageBreak as RLPageBreak
from reportlab.platypus import Paragraph as RLParagraph
from reportlab.platypus import SimpleDocTemplate
from reportlab.platypus import Spacer as RLSpacer
from reportlab.platypus import Table as RLTable
from reportlab.platypus import TableStyle

from apps.core.timeutils import now

from .blocks import PageBreak, Paragraph, Signatures, Spacer, Subtitle, Table, Title

NAVY = colors.HexColor("#0B1F4D")
GOLD = colors.HexColor("#C9A227")
RED = colors.HexColor("#B91C1C")

CONTENT_TYPES = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}


def _footer(canvas, doc):
    canvas.saveState()
    canvas.setFont("Helvetica", 7)
    canvas.setFillColor(colors.grey)
    canvas.drawString(1.5 * cm, 1 * cm, f"FutOwl · Generado el {now():%d/%m/%Y %H:%M} (VET, UTC-4)")
    canvas.drawRightString(A4[0] - 1.5 * cm, 1 * cm, f"Página {doc.page}")
    canvas.setStrokeColor(GOLD)
    canvas.line(1.5 * cm, A4[1] - 1.2 * cm, A4[0] - 1.5 * cm, A4[1] - 1.2 * cm)
    canvas.restoreState()


def render_pdf(blocks, title: str) -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, title=title, author="FutOwl",
                            leftMargin=1.5 * cm, rightMargin=1.5 * cm, topMargin=1.6 * cm, bottomMargin=1.6 * cm)
    styles = getSampleStyleSheet()
    s_title = ParagraphStyle("t", parent=styles["Title"], textColor=NAVY, fontSize=16, spaceAfter=6)
    s_sub = ParagraphStyle("s", parent=styles["Heading3"], textColor=RED, spaceBefore=6, spaceAfter=4)
    s_par = ParagraphStyle("p", parent=styles["BodyText"], fontSize=9.5, leading=13)
    s_cell = ParagraphStyle("c", parent=styles["BodyText"], fontSize=8, leading=10)
    story = []
    width = A4[0] - 3 * cm
    for block in blocks:
        if isinstance(block, Title):
            story.append(RLParagraph(escape(block.text), s_title))
        elif isinstance(block, Subtitle):
            story.append(RLParagraph(escape(block.text), s_sub))
        elif isinstance(block, Paragraph):
            text = escape(block.text)
            text = f"<b>{text}</b>" if block.bold else text
            story.append(RLParagraph(text.replace("\n", "<br/>"), s_par))
        elif isinstance(block, Spacer):
            story.append(RLSpacer(1, block.height))
        elif isinstance(block, PageBreak):
            story.append(RLPageBreak())
        elif isinstance(block, Table):
            widths = block.widths or [1] * len(block.headers)
            total = sum(widths)
            col_widths = [width * w / total for w in widths]
            data = [[RLParagraph(f"<b>{escape(h)}</b>", s_cell) for h in block.headers]]
            data += [[RLParagraph(escape(str(c)), s_cell) for c in row] for row in block.rows]
            table = RLTable(data, colWidths=col_widths, repeatRows=1)
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), NAVY),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#94A3B8")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F5F9")]),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ]))
            story.append(table)
        elif isinstance(block, Signatures):
            cells = [["_" * 28 for _ in block.labels], block.labels]
            sig = RLTable(cells, colWidths=[width / len(block.labels)] * len(block.labels))
            sig.setStyle(TableStyle([("ALIGN", (0, 0), (-1, -1), "CENTER"), ("FONTSIZE", (0, 0), (-1, -1), 8),
                                     ("TOPPADDING", (0, 0), (-1, 0), 28)]))
            story.append(sig)
    doc.build(story, onFirstPage=_footer, onLaterPages=_footer)
    return buffer.getvalue()


def render_docx(blocks, title: str) -> bytes:
    document = Document()
    document.core_properties.title = title
    document.core_properties.author = "FutOwl"
    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10)
    for block in blocks:
        if isinstance(block, Title):
            heading = document.add_heading(block.text, level=1)
            heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in heading.runs:
                run.font.color.rgb = RGBColor(0x0B, 0x1F, 0x4D)
        elif isinstance(block, Subtitle):
            heading = document.add_heading(block.text, level=3)
            for run in heading.runs:
                run.font.color.rgb = RGBColor(0xB9, 0x1C, 0x1C)
        elif isinstance(block, Paragraph):
            paragraph = document.add_paragraph()
            run = paragraph.add_run(block.text)
            run.bold = block.bold
        elif isinstance(block, Spacer):
            document.add_paragraph()
        elif isinstance(block, PageBreak):
            document.add_page_break()
        elif isinstance(block, Table):
            table = document.add_table(rows=1, cols=len(block.headers))
            table.style = "Table Grid"
            for cell, header in zip(table.rows[0].cells, block.headers):
                cell.text = ""
                cell.paragraphs[0].add_run(header).bold = True
            for row in block.rows:
                cells = table.add_row().cells
                for cell, value in zip(cells, row):
                    cell.text = str(value)
        elif isinstance(block, Signatures):
            table = document.add_table(rows=2, cols=len(block.labels))
            for i, label in enumerate(block.labels):
                table.rows[0].cells[i].text = "\n\n" + "_" * 24
                table.rows[1].cells[i].text = label
    footer = document.sections[0].footer.paragraphs[0]
    footer.text = f"FutOwl · Generado el {now():%d/%m/%Y %H:%M} (VET, UTC-4)"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def render(blocks, fmt: str, title: str) -> bytes:
    if fmt == "docx":
        return render_docx(blocks, title)
    return render_pdf(blocks, title)
